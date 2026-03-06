"""
paper_trader.py — Excel-based paper trading manager for the S4 strategy.

All portfolio state is stored in 's4_paper_trading.xlsx'.
Peak drawdown tracking is stored in '.paper_peak.json'.
"""

import datetime
import json
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from s4_engine import CAPITAL, COST_PER_SIDE

PEAK_FILE = '.paper_peak.json'

# ---------------------------------------------------------------------------
# Indian number formatting
# ---------------------------------------------------------------------------

def format_indian(num):
    """Format number in Indian style: 20,45,231"""
    s = str(abs(int(num)))
    if len(s) <= 3:
        result = s
    else:
        result = s[-3:]
        s = s[:-3]
        while s:
            result = s[-2:] + ',' + result
            s = s[:-2]
    sign = '-' if num < 0 else ''
    return sign + result


# ---------------------------------------------------------------------------
# Excel styling helpers
# ---------------------------------------------------------------------------

HEADER_FILL   = PatternFill('solid', fgColor='1F2937')
HEADER_FONT   = Font(color='FFFFFF', bold=True)
HEADER_ALIGN  = Alignment(horizontal='center', vertical='center')
BORDER_SIDE   = Side(style='thin', color='D1D5DB')
THIN_BORDER   = Border(
    left=BORDER_SIDE, right=BORDER_SIDE,
    top=BORDER_SIDE,  bottom=BORDER_SIDE,
)

def _style_header_row(ws, headers, row=1):
    for col_idx, hdr in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=hdr)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 22

def _auto_width(ws, min_width=10):
    for col in ws.columns:
        max_len = min_width
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or '')))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 30)


# ---------------------------------------------------------------------------
# 1. create_workbook
# ---------------------------------------------------------------------------

SHEET_HEADERS = {
    'Dashboard': ['Key', 'Value'],
    'Portfolio': [
        'Ticker', 'Qty', 'Avg Price', 'Current Price',
        'Value', 'P&L', 'P&L %', 'Day P&L', 'Status',
    ],
    'Trades': [
        'Date', 'Ticker', 'Action', 'Qty', 'Price',
        'Value', 'Commission', 'Regime',
    ],
    'Equity': ['Date', 'Portfolio Value', 'Cash', 'Holdings Value', 'Drawdown %'],
    'Signals': [
        'Date', 'Regime', 'Breadth %', 'Vol Fast %', 'Vol Slow %',
        'Vol Scale', 'Total Weight', 'N Above SMA', 'N Positive Mom',
    ],
}


def create_workbook(filepath) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    for name, headers in SHEET_HEADERS.items():
        ws = wb.create_sheet(name)
        _style_header_row(ws, headers)
        _auto_width(ws)

    # Seed Equity sheet with starting capital row
    ws_eq = wb['Equity']
    ws_eq.append([
        datetime.date.today().isoformat(),
        CAPITAL, CAPITAL, 0, 0.0,
    ])

    wb.save(filepath)
    return wb


# ---------------------------------------------------------------------------
# 2. load_or_create
# ---------------------------------------------------------------------------

def load_or_create(filepath) -> Workbook:
    if os.path.exists(filepath):
        try:
            return load_workbook(filepath)
        except Exception:
            pass
    return create_workbook(filepath)


# ---------------------------------------------------------------------------
# 3. load_portfolio
# ---------------------------------------------------------------------------

def load_portfolio(filepath):
    """
    Returns (holdings_dict, cash).
    holdings_dict: {ticker: {qty, avg_price, current_price}}
    """
    wb = load_or_create(filepath)
    holdings = {}

    ws_port = wb['Portfolio']
    for row in ws_port.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        ticker, qty, avg_price, current_price = row[0], row[1], row[2], row[3]
        if ticker and qty:
            holdings[str(ticker)] = {
                'qty':           int(qty or 0),
                'avg_price':     float(avg_price or 0),
                'current_price': float(current_price or avg_price or 0),
            }

    # Cash: last row of Equity sheet
    ws_eq = wb['Equity']
    cash = CAPITAL
    for row in ws_eq.iter_rows(min_row=2, values_only=True):
        if row and row[2] is not None:
            try:
                cash = float(row[2])
            except (TypeError, ValueError):
                pass

    return holdings, cash


# ---------------------------------------------------------------------------
# 4. compute_trades
# ---------------------------------------------------------------------------

def compute_trades(holdings, cash, signals):
    """
    Compare current holdings to target signals.

    Returns (sells, buys, portfolio_value).
    """
    target = signals.get('holdings', {})
    live_prices = {t: d['price'] for t, d in target.items()}

    # Estimate portfolio value (mark holdings to target prices or kept prices)
    holdings_value = sum(
        h['qty'] * live_prices.get(t, h['current_price'])
        for t, h in holdings.items()
    )
    portfolio_value = holdings_value + cash
    invest_capital = portfolio_value * signals.get('total_weight', 1.0)

    sells = []
    buys  = []
    BUFFER = 5  # share buffer to avoid tiny trades

    # --- Sells / Reduces ---
    for ticker, h in holdings.items():
        price = live_prices.get(ticker, h['current_price'])
        if price <= 0:
            continue
        if ticker not in target:
            # Full exit
            sells.append({
                'ticker': ticker,
                'action': 'EXIT',
                'qty':    h['qty'],
                'price':  round(price, 2),
            })
        else:
            # Check if qty needs to reduce
            target_w   = target[ticker]['weight']
            target_qty = int(invest_capital * target_w / price)
            diff = h['qty'] - target_qty
            if diff > BUFFER:
                sells.append({
                    'ticker': ticker,
                    'action': 'REDUCE',
                    'qty':    diff,
                    'price':  round(price, 2),
                })

    # --- Buys / Adds ---
    for ticker, t in target.items():
        price = t['price']
        if price <= 0:
            continue
        target_qty = int(invest_capital * t['weight'] / price)
        current_qty = holdings.get(ticker, {}).get('qty', 0)
        diff = target_qty - current_qty
        if diff > BUFFER:
            buys.append({
                'ticker': ticker,
                'action': 'NEW' if ticker not in holdings else 'ADD',
                'qty':    diff,
                'price':  round(price, 2),
            })

    return sells, buys, portfolio_value


# ---------------------------------------------------------------------------
# 5. execute_trades
# ---------------------------------------------------------------------------

def execute_trades(filepath, holdings, cash, sells, buys, signals):
    """
    Execute sell then buy orders, update all Excel sheets.

    Returns (updated_holdings, cash, portfolio_value).
    """
    wb = load_or_create(filepath)
    today = datetime.date.today().isoformat()
    regime = signals.get('regime', 'NEUTRAL')

    updated = {t: dict(h) for t, h in holdings.items()}

    # ---- Process Sells ----
    for trade in sells:
        ticker = trade['ticker']
        qty    = trade['qty']
        price  = trade['price']
        commission = price * qty * COST_PER_SIDE
        proceeds = price * qty - commission
        cash += proceeds

        if ticker in updated:
            updated[ticker]['qty'] -= qty
            if updated[ticker]['qty'] <= 0:
                del updated[ticker]

        _append_trade(wb, today, ticker, 'SELL', qty, price, commission, regime)

    # ---- Process Buys ----
    for trade in buys:
        ticker = trade['ticker']
        qty    = trade['qty']
        price  = trade['price']
        commission = price * qty * COST_PER_SIDE
        cost = price * qty + commission

        if cost > cash:
            # Reduce qty to fit available cash
            qty = int((cash * 0.99) / (price * (1 + COST_PER_SIDE)))
            if qty <= 0:
                continue
            commission = price * qty * COST_PER_SIDE
            cost = price * qty + commission

        cash -= cost

        if ticker in updated:
            prev_qty   = updated[ticker]['qty']
            prev_avg   = updated[ticker]['avg_price']
            new_qty    = prev_qty + qty
            new_avg    = (prev_qty * prev_avg + qty * price) / new_qty
            updated[ticker]['qty']       = new_qty
            updated[ticker]['avg_price'] = round(new_avg, 2)
        else:
            updated[ticker] = {
                'qty':           qty,
                'avg_price':     round(price, 2),
                'current_price': round(price, 2),
            }

        _append_trade(wb, today, ticker, 'BUY', qty, price, commission, regime)

    # ---- Recalculate portfolio value ----
    target_prices = {t: d['price'] for t, d in signals.get('holdings', {}).items()}
    holdings_value = sum(
        h['qty'] * target_prices.get(t, h.get('current_price', h['avg_price']))
        for t, h in updated.items()
    )
    portfolio_value = holdings_value + cash

    # ---- Track peak ----
    peak = _load_peak()
    if portfolio_value > peak:
        peak = portfolio_value
        _save_peak(peak)

    # ---- Update sheets ----
    _overwrite_portfolio(wb, updated, target_prices)
    _append_equity(wb, today, portfolio_value, cash, holdings_value, peak)
    _append_signal(wb, today, signals)
    _write_dashboard(wb, portfolio_value, cash, holdings_value, peak, regime, updated)

    wb.save(filepath)
    return updated, cash, portfolio_value


# ---------------------------------------------------------------------------
# 6. update_prices
# ---------------------------------------------------------------------------

def update_prices(filepath, live_prices) -> dict:
    """
    Update Portfolio sheet with new current prices.
    Returns summary dict.
    """
    wb = load_or_create(filepath)
    ws = wb['Portfolio']

    total_value  = 0.0
    total_pnl    = 0.0
    day_pnl      = 0.0
    cash         = 0.0

    # Read previous prices for day P&L
    prev_prices = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            prev_prices[row[0]] = float(row[3] or row[2] or 0)

    for row in ws.iter_rows(min_row=2):
        ticker = row[0].value
        if not ticker:
            continue
        qty = int(row[1].value or 0)
        avg = float(row[2].value or 0)
        old_price = float(row[3].value or avg)
        new_price = live_prices.get(ticker, old_price)

        value    = qty * new_price
        pnl      = (new_price - avg) * qty
        pnl_pct  = ((new_price / avg) - 1) * 100 if avg else 0
        d_pnl    = (new_price - old_price) * qty

        row[3].value = round(new_price, 2)
        row[4].value = round(value, 2)
        row[5].value = round(pnl, 2)
        row[6].value = round(pnl_pct, 2)
        row[7].value = round(d_pnl, 2)

        total_value += value
        total_pnl   += pnl
        day_pnl     += d_pnl

    # Cash
    ws_eq = wb['Equity']
    for row in ws_eq.iter_rows(min_row=2, values_only=True):
        if row and row[2] is not None:
            try:
                cash = float(row[2])
            except (TypeError, ValueError):
                pass

    portfolio_value = total_value + cash
    total_return    = ((portfolio_value / CAPITAL) - 1) * 100 if CAPITAL else 0

    peak = _load_peak()
    if portfolio_value > peak:
        peak = portfolio_value
        _save_peak(peak)
    drawdown = ((portfolio_value / peak) - 1) * 100 if peak else 0

    wb.save(filepath)
    return {
        'portfolio_value': round(portfolio_value, 2),
        'total_pnl':       round(total_pnl, 2),
        'total_return':    round(total_return, 2),
        'drawdown':        round(drawdown, 2),
        'day_pnl':         round(day_pnl, 2),
    }


# ---------------------------------------------------------------------------
# 7. get_equity_history
# ---------------------------------------------------------------------------

def get_equity_history(filepath):
    """Returns (dates_list, values_list) for charting."""
    wb = load_or_create(filepath)
    ws = wb['Equity']
    dates  = []
    values = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] and row[1]:
            dates.append(str(row[0]))
            try:
                values.append(float(row[1]))
            except (TypeError, ValueError):
                values.append(0.0)
    return dates, values


# ---------------------------------------------------------------------------
# 8. get_trade_history
# ---------------------------------------------------------------------------

def get_trade_history(filepath) -> list:
    """Returns list of trade dicts."""
    wb = load_or_create(filepath)
    ws = wb['Trades']
    trades = []
    headers = SHEET_HEADERS['Trades']
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            trades.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    return trades


# ---------------------------------------------------------------------------
# 9. get_stats
# ---------------------------------------------------------------------------

def get_stats(filepath) -> dict:
    """Calculate performance stats from Equity sheet."""
    dates, values = get_equity_history(filepath)
    if len(values) < 2:
        return {
            'cagr': 0.0, 'total_return': 0.0,
            'max_dd': 0.0, 'win_rate': 0.0,
            'best_day': 0.0, 'worst_day': 0.0,
        }

    series = pd.Series(values)
    rets   = series.pct_change().dropna()

    # Total return
    total_return = (series.iloc[-1] / series.iloc[0] - 1) * 100

    # CAGR
    try:
        start_dt = datetime.date.fromisoformat(str(dates[0]))
        end_dt   = datetime.date.fromisoformat(str(dates[-1]))
        years    = max((end_dt - start_dt).days / 365.25, 1 / 365)
        cagr     = ((series.iloc[-1] / series.iloc[0]) ** (1 / years) - 1) * 100
    except Exception:
        cagr = total_return

    # Max drawdown
    roll_max = series.cummax()
    dd       = (series / roll_max - 1) * 100
    max_dd   = float(dd.min())

    win_rate  = float((rets > 0).mean() * 100) if len(rets) else 0.0
    best_day  = float(rets.max() * 100) if len(rets) else 0.0
    worst_day = float(rets.min() * 100) if len(rets) else 0.0

    return {
        'cagr':         round(cagr, 2),
        'total_return': round(total_return, 2),
        'max_dd':       round(max_dd, 2),
        'win_rate':     round(win_rate, 1),
        'best_day':     round(best_day, 2),
        'worst_day':    round(worst_day, 2),
    }


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _append_trade(wb, date, ticker, action, qty, price, commission, regime):
    ws = wb['Trades']
    ws.append([date, ticker, action, qty, round(price, 2),
               round(price * qty, 2), round(commission, 2), regime])


def _overwrite_portfolio(wb, holdings, target_prices):
    ws = wb['Portfolio']
    # Clear existing data rows
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.value = None
    # Write new rows
    for row_idx, (ticker, h) in enumerate(holdings.items(), start=2):
        qty   = h['qty']
        avg   = h['avg_price']
        cur   = target_prices.get(ticker, h.get('current_price', avg))
        val   = qty * cur
        pnl   = (cur - avg) * qty
        pnl_p = ((cur / avg) - 1) * 100 if avg else 0
        status = '▲ Profit' if pnl >= 0 else '▼ Loss'
        ws.append([ticker, qty, round(avg, 2), round(cur, 2),
                   round(val, 2), round(pnl, 2), round(pnl_p, 2), 0.0, status])


def _append_equity(wb, date, portfolio_value, cash, holdings_value, peak):
    ws = wb['Equity']
    drawdown = ((portfolio_value / peak) - 1) * 100 if peak else 0
    ws.append([date, round(portfolio_value, 2), round(cash, 2),
               round(holdings_value, 2), round(drawdown, 2)])


def _append_signal(wb, date, signals):
    ws = wb['Signals']
    ws.append([
        date,
        signals.get('regime'),
        signals.get('breadth'),
        signals.get('vol_fast'),
        signals.get('vol_slow'),
        signals.get('vol_scale'),
        signals.get('total_weight'),
        signals.get('n_above_sma'),
        signals.get('n_positive_mom'),
    ])


def _write_dashboard(wb, portfolio_value, cash, holdings_value, peak, regime, holdings):
    ws = wb['Dashboard']
    # Clear
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.value = None
    rows = [
        ('Portfolio Value', f'Rs {format_indian(portfolio_value)}'),
        ('Cash',           f'Rs {format_indian(cash)}'),
        ('Holdings Value', f'Rs {format_indian(holdings_value)}'),
        ('Peak Value',     f'Rs {format_indian(peak)}'),
        ('P&L',            f'Rs {format_indian(portfolio_value - CAPITAL)}'),
        ('Return %',       f'{round((portfolio_value / CAPITAL - 1) * 100, 2)} %'),
        ('Regime',         regime),
        ('# Holdings',     len(holdings)),
        ('Last Updated',   datetime.datetime.now().strftime('%Y-%m-%d %H:%M')),
    ]
    for r in rows:
        ws.append(list(r))


def _load_peak() -> float:
    if os.path.exists(PEAK_FILE):
        try:
            with open(PEAK_FILE) as f:
                return float(json.load(f).get('peak', CAPITAL))
        except Exception:
            pass
    return float(CAPITAL)


def _save_peak(peak: float):
    with open(PEAK_FILE, 'w') as f:
        json.dump({'peak': peak}, f)
